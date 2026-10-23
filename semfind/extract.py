"""Estrazione di testo e immagini dai file, e divisione in pezzi."""
import json
import re

import pymupdf
from PIL import Image

pymupdf.TOOLS.mupdf_display_errors(False)

Chunk = tuple[int | None, str]  # (pagina, testo)


def split(text: str, size: int, overlap: int) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= size:
        return [text] if text else []
    out, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            # taglia su un confine naturale se ce n'è uno nell'ultimo terzo
            cut = max(text.rfind("\n\n", start, end), text.rfind(". ", start, end))
            if cut > start + size * 2 // 3:
                end = cut + 1
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = end - overlap
    return [c for c in out if c]


def _pages_pdf(path: str) -> tuple[list[tuple[int, str]], int]:
    with pymupdf.open(path) as doc:
        return [(i + 1, p.get_text()) for i, p in enumerate(doc)], doc.page_count


def text_chunks(path: str, ext: str, cfg) -> tuple[list[Chunk], bool]:
    """Ritorna (pezzi, scanned). scanned=True se è un PDF senza testo estraibile."""
    size, ov = cfg.chunk_chars, cfg.chunk_overlap
    chunks: list[Chunk] = []
    if ext in ("pdf", "epub"):
        pages, n = _pages_pdf(path)
        total = sum(len(t.strip()) for _, t in pages)
        if ext == "pdf" and n and total / n < 50:
            return [], True
        for page, t in pages:
            if len(t.strip()) < 40:
                continue
            chunks += [(page if ext == "pdf" else None, c) for c in split(t, size, ov)]
    elif ext == "docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        chunks = [(None, c) for c in split("\n".join(parts), size, ov)]
    elif ext == "pptx":
        import pptx
        prs = pptx.Presentation(path)
        for i, slide in enumerate(prs.slides, 1):
            t = "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            chunks += [(i, c) for c in split(t, size, ov)]
    elif ext == "ipynb":
        nb = json.loads(open(path, encoding="utf-8", errors="replace").read())
        t = "\n\n".join("".join(c.get("source", "")) for c in nb.get("cells", []))
        chunks = [(None, c) for c in split(t, size, ov)]
    else:
        with open(path, encoding="utf-8", errors="replace") as f:
            t = f.read()
        if "\x00" in t[:4096]:
            return [], False
        chunks = [(None, c) for c in split(t, size, ov)]
    return chunks[: cfg.max_chunks_per_file], False


def load_image(path: str, cfg, max_side: int = 1024) -> Image.Image | None:
    im = Image.open(path)
    if min(im.size) < cfg.min_image_side:
        return None
    im.draft("RGB", (max_side, max_side))  # decodifica JPEG ridotta, molto più veloce
    im = im.convert("RGB")
    im.thumbnail((max_side, max_side))
    return im


def pdf_page_images(path: str, cfg) -> list[tuple[int, Image.Image]]:
    out = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc):
            if i >= cfg.scanned_pdf_max_pages:
                break
            zoom = 1024 / max(page.rect.width, page.rect.height)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csRGB)
            out.append((i + 1, Image.frombytes("RGB", (pix.width, pix.height), pix.samples)))
    return out
